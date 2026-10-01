"""بيطلّع ملف Excel فيه شيتين: Old و New، لنفس العملاء.

Old: العلامات الأصلية (النقاط المتبقية من كل بُعد) والـ CEI القديم.
New: نفس العملاء بعد تطبيق الأوزان الجديدة على كل بُعد، والـ CEI الجديد.

الأوزان الجديدة مكتوبة بالسطر 2 من شيت New. كل خانة بشيت New معادلة:
    النقاط الجديدة = (علامة العميل ÷ الوزن القديم) × الوزن الجديد
    CEI الجديد = مجموع النقاط الجديدة للأبعاد الثمانية
يعني إذا غيّرت أي وزن بالسطر 2 (جوا Excel)، كل السكورات والـ CEI بيتحدثوا لحالهم.
وإذا بدك تغيّرها من Python، غيّر NEW_WEIGHTS تحت وشغّل الملف من جديد.

الاستخدام:
    python cei_old_new.py                      # بيستخدم FILE_PATH و NEW_WEIGHTS اللي تحت
    python cei_old_new.py <merged.csv> [output.xlsx]

أو من cei_advice_model.py بعد ما يطلع final:
    from cei_old_new import write_old_new
    write_old_new(data, dict(zip(COLUMNS, final)), INPUT.with_name(INPUT.stem + "_old_new.xlsx"))

Requires: pandas, numpy, openpyxl
"""
from pathlib import Path
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FILE_PATH = r"C:\Users\Dell\Downloads\Pro\Pro\merged_cei_20261001_094541.csv"
SERIAL_COLUMN = "SERIAL_NUMBER"
ADVICE_COLUMN = "Repair Advice"

OLD_WEIGHTS = {
    "Service Score": 30, "Wi-Fi Score": 20, "Rate Score": 5, "Stability Score": 15,
    "STA Score": 10, "Gateway Score": 10, "ODN Score": 5, "OLT Score": 5,
}
# غيّر هون وشغّل من جديد، أو غيّر السطر 2 بشيت New جوا Excel
NEW_WEIGHTS = {
    "Service Score": 2, "Wi-Fi Score": 39, "Rate Score": 15, "Stability Score": 2,
    "STA Score": 25, "Gateway Score": 13, "ODN Score": 2, "OLT Score": 2,
}
COLUMNS = list(OLD_WEIGHTS)

FONT = "Arial"
HEAD = dict(font=Font(name=FONT, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor="000000"),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
WEIGHT = dict(font=Font(name=FONT, bold=True), fill=PatternFill("solid", fgColor="D9D9D9"),
              border=Border(*(Side(style="thin", color="000000"),) * 4), alignment=Alignment(horizontal="center"))
LABEL = dict(font=Font(name=FONT, bold=True))
NOTE = dict(font=Font(name=FONT, italic=True, size=9, color="595959"))


def styled(ws, value, style, fmt=None):
    c = WriteOnlyCell(ws, value=value)
    for k, v in style.items():
        setattr(c, k, v)
    if fmt:
        c.number_format = fmt
    return c


def write_old_new(data, new_weights, output_path):
    """data: الداتا الأصلية (فيها SERIAL_NUMBER و Repair Advice والأبعاد الثمانية)."""
    new_weights = {c: float(new_weights[c]) for c in COLUMNS}
    if abs(sum(new_weights.values()) - 100) > 1e-9:
        print(f"Warning: new weights sum to {sum(new_weights.values())}, not 100.")
    scores = data[COLUMNS].apply(pd.to_numeric, errors="raise")
    for col in COLUMNS:
        if not scores[col].between(0, OLD_WEIGHTS[col]).all():
            raise ValueError(f"{col}: scores must be between 0 and {OLD_WEIGHTS[col]}")
    n = len(data)
    first, last = 5, 4 + n            # الداتا من السطر 5
    score_cols = [get_column_letter(3 + i) for i in range(len(COLUMNS))]   # C..J
    cei_col = get_column_letter(3 + len(COLUMNS))                           # K

    wb = Workbook(write_only=True)
    for name, weights, title in [("Old", OLD_WEIGHTS, "Old weights (original scores)"),
                                 ("New", new_weights, "New weights: edit row 2 and every score and CEI updates")]:
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "C5"
        widths = [20, 45 if name == "Old" else 18] + [13] * len(COLUMNS) + [11]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        weight_cells = [styled(ws, weights[c], WEIGHT, "0.##") for c in COLUMNS]
        total = styled(ws, f"=SUM(C2:{score_cols[-1]}2)", WEIGHT, "0.##")
        ws.append([styled(ws, title, LABEL)])
        ws.append([styled(ws, "Weight (full score)", LABEL), None] + weight_cells + [total])
        ws.append([styled(ws, "Old: score as in the file.  New: =(Old score ÷ Old weight) × New weight.  "
                              "CEI = sum of the eight scores.", NOTE)])
        second = ADVICE_COLUMN if name == "Old" else "Has Repair Advice"
        ws.append([styled(ws, h, HEAD) for h in [SERIAL_COLUMN, second] + COLUMNS + ["CEI"]])

        serials = data[SERIAL_COLUMN].tolist()
        text = data[ADVICE_COLUMN].fillna("")
        advice = (text.tolist() if name == "Old"   # النص الكامل بشيت Old بس، عشان حجم الملف
                  else np.where(text.str.contains("ADVICE_"), "Yes", "No").tolist())
        values = scores.to_numpy()
        for r in range(n):
            row_no = first + r
            if name == "Old":
                cells = [float(v) if v % 1 else int(v) for v in values[r]]
            else:
                cells = []
                for L in score_cols:
                    c = WriteOnlyCell(ws, value=f"=Old!{L}{row_no}/Old!{L}$2*{L}$2")
                    c.number_format = "0.00"
                    cells.append(c)
            cei = WriteOnlyCell(ws, value=f"=SUM(C{row_no}:{score_cols[-1]}{row_no})")
            cei.number_format = "0.00"
            ws.append([serials[r], advice[r]] + cells + [cei])

    wb.calculation.fullCalcOnLoad = True
    wb.save(output_path)

    # نفس الحساب بـ Python للتأكد ولطباعة ملخص
    old_w = np.array([OLD_WEIGHTS[c] for c in COLUMNS], float)
    new_w = np.array([new_weights[c] for c in COLUMNS], float)
    ratio = scores.to_numpy(float) / old_w
    cei_old, cei_new = ratio @ old_w, ratio @ new_w
    print(f"Saved: {output_path}")
    print(f"  Customers: {n:,}")
    print(f"  Average CEI: old {cei_old.mean():.2f}  new {cei_new.mean():.2f}")
    print(f"  CEI < 80: old {(cei_old < 80 - 1e-9).sum():,}  new {(cei_new < 80 - 1e-9).sum():,}")
    return cei_old, cei_new


def read_csv(path):
    try:
        data = pd.read_csv(path, dtype=str, encoding="utf-8-sig")
    except UnicodeDecodeError:
        data = pd.read_csv(path, dtype=str, encoding="cp1256")
    data.columns = [" ".join(str(c).split()) for c in data.columns]
    return data


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(FILE_PATH)
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem + "_old_new.xlsx")
    write_old_new(read_csv(src), NEW_WEIGHTS, out)
