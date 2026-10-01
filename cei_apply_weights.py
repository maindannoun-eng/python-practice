"""يطبّق أوزان الـ CEI الجديدة على كل العملاء، ويطلّع ملف Excel جاهز للعرض والمقارنة.

بيقرأ:
  1. الملف المدموج (نفس الملف اللي اشتغل عليه cei_advice_model.py).
  2. ملف النتائج تبع الموديل (*_advice_model_result.xlsx) عشان ياخد الأوزان النهائية
     ونتائج عينة الاختبار. إذا مش موجود، بيستخدم الأوزان المكتوبة تحت بـ FALLBACK_NEW_WEIGHTS.

بيطلّع ملف Excel فيه:
  Overview             شرح الملف وكيف انحسب كل رقم
  Summary              القديم مقابل الجديد: كل العملاء + عينة الاختبار
  Weights              الأوزان القديمة والجديدة
  Dimension_Comparison لكل بُعد: الوزن، متوسط النقاط، ووين بينخصم
  Band_Distribution    توزيع العملاء على فئات الـ CEI (مع رسمة)
  Band_Transition      انتقال العملاء من فئة لفئة
  Status_Transition    انتقال الحالة (مكتشف / مخفي / إنذار خاطئ / سليم)
  Per_Advice           لكل نصيحة: الكشف قبل وبعد
  Per_Family           نفس الشي لكل عائلة نصايح
  Customer_Data        كل عميل: السكورات القديمة والجديدة لكل بُعد، والـ CEI القديم والجديد
                       (جدول Excel جاهز للـ Pivot Table)

كل الأرقام بشيتات المقارنة معادلات (COUNTIFS / AVERAGEIFS) على شيت Customer_Data،
يعني بتتحدث لحالها وبتقدر تتتبع كل رقم من وين جاي.

Requires: pandas, numpy, openpyxl
"""
from datetime import date
from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

# ---------- المسارات ----------
FILE_PATH = r"C:\Users\Dell\Downloads\Pro\Pro\merged_cei_20261001_094541.csv"
INPUT = Path(FILE_PATH)
RESULT_PATH = INPUT.with_name(INPUT.stem + "_advice_model_result.xlsx")  # ناتج cei_advice_model.py
OUTPUT_PATH = INPUT.with_name(INPUT.stem + "_new_weights_applied.xlsx")

# ---------- الإعدادات (نفس cei_advice_model.py) ----------
SERIAL_COLUMN = "SERIAL_NUMBER"
ADVICE_COLUMN = "Repair Advice"
ADVICE_REGEX = r"ADVICE_[A-Z0-9_]+"
CEI_THRESHOLD = 80
TEST_SHARE = 0.3
SEED = 42
BAND_SIZE = 5

OLD_WEIGHTS = {
    "Service Score": 30, "Wi-Fi Score": 20, "Rate Score": 5, "Stability Score": 15,
    "STA Score": 10, "Gateway Score": 10, "ODN Score": 5, "OLT Score": 5,
}
# بتُستخدم بس إذا ملف النتائج مش موجود
FALLBACK_NEW_WEIGHTS = {
    "Service Score": 2, "Wi-Fi Score": 39, "Rate Score": 15, "Stability Score": 2,
    "STA Score": 25, "Gateway Score": 13, "ODN Score": 2, "OLT Score": 2,
}

COLUMNS = list(OLD_WEIGHTS)
SHORT = [c.replace(" Score", "") for c in COLUMNS]


def advice_family(code):
    """عائلة النصيحة من اسمها (للتجميع بس)."""
    name = code.replace("ADVICE_", "")
    if name in ("HIGH_MEMORY_USAGE", "HIGH_CPU_USAGE"):
        return "Gateway resources"
    if name.startswith(("OLT_", "ONT_")):
        return "OLT / ONT access"
    if "OPTICAL" in name and not name.startswith("AP_"):
        return "Gateway optical"
    if name.startswith(("AP_", "DUAL_BAND_AP_", "EXTERN_AP_")):
        return "AP / FTTR"
    return "Wi-Fi"


FAMILIES = ["Wi-Fi", "AP / FTTR", "Gateway optical", "Gateway resources", "OLT / ONT access"]


# ======================== الحساب ========================

def load_weights():
    if RESULT_PATH.exists():
        table = pd.read_excel(RESULT_PATH, sheet_name="Final_Weights", index_col=0)
        new = {c: int(table.loc[c, "Final"]) for c in COLUMNS}
        source = f"{RESULT_PATH.name}, sheet Final_Weights"
    else:
        new = dict(FALLBACK_NEW_WEIGHTS)
        source = "FALLBACK_NEW_WEIGHTS in this script"
    if sum(new.values()) != 100:
        raise ValueError(f"New weights must sum to 100, got {sum(new.values())}")
    return new, source


def load_model_summary():
    if not RESULT_PATH.exists():
        return None
    try:
        return pd.read_excel(RESULT_PATH, sheet_name="Summary", index_col=0)
    except ValueError:
        return None


def band_label(cei):
    low = np.minimum(np.floor(cei / BAND_SIZE) * BAND_SIZE, 100 - BAND_SIZE).astype(int)
    return [f"{a:02d}-{a + BAND_SIZE}" for a in low]


def build_data(new_weights):
    print("Reading the input file...", flush=True)
    try:
        raw = pd.read_csv(FILE_PATH, dtype=str, encoding="utf-8-sig")
    except UnicodeDecodeError:
        raw = pd.read_csv(FILE_PATH, dtype=str, encoding="cp1256")
    raw.columns = [" ".join(str(c).split()) for c in raw.columns]
    missing = [c for c in COLUMNS + [SERIAL_COLUMN, ADVICE_COLUMN] if c not in raw.columns]
    if missing:
        raise ValueError(f"Missing columns: {missing}")
    scores = raw[COLUMNS].apply(pd.to_numeric, errors="raise")
    for col in COLUMNS:
        if not scores[col].between(0, OLD_WEIGHTS[col]).all():
            raise ValueError(f"{col}: scores must be between 0 and {OLD_WEIGHTS[col]}")

    codes = (raw[ADVICE_COLUMN].fillna("").str.upper()
             .str.findall(ADVICE_REGEX).apply(lambda f: sorted(set(f), key=f.index)))
    old_w = np.array([OLD_WEIGHTS[c] for c in COLUMNS], dtype=float)
    new_w = np.array([new_weights[c] for c in COLUMNS], dtype=float)
    achieved = scores.to_numpy(float) / old_w           # النسبة اللي حققها العميل من كل بُعد
    new_points = np.round(achieved * new_w, 4)            # نقاط كل بُعد بالأوزان الجديدة
    cei_old = np.round(achieved @ old_w, 6)
    cei_new = np.round(achieved @ new_w, 6)               # نفس حساب cei_advice_model.py
    has = codes.str.len().to_numpy() > 0
    is_test = np.random.default_rng(SEED).random(len(raw)) < TEST_SHARE  # نفس تقسيم الموديل

    def status(cei):
        low = cei < CEI_THRESHOLD
        return np.select([has & low, has & ~low, ~has & low],
                         ["Detected", "Hidden", "False alarm"], "OK")

    low_old, low_new = cei_old < CEI_THRESHOLD, cei_new < CEI_THRESHOLD
    out = pd.DataFrame({
        "Serial_Number": raw[SERIAL_COLUMN],
        "Advice_Codes": codes.apply(lambda c: ", ".join(x.replace("ADVICE_", "") for x in c)),
        "Advice_Count": codes.str.len(),
        "Has_Repair_Advice": np.where(has, "Yes", "No"),
        "Main_Advice": codes.apply(lambda c: c[0].replace("ADVICE_", "") if c else "None"),
        "Advice_Families": codes.apply(lambda c: ", ".join(sorted({advice_family(x) for x in c},
                                                                  key=FAMILIES.index))),
        "Sample": np.where(is_test, "Test", "Train"),
    })
    for i, s in enumerate(SHORT):
        out[f"{s}_Old"] = scores[COLUMNS[i]].to_numpy()
    for i, s in enumerate(SHORT):
        out[f"{s}_New"] = new_points[:, i]
    out["CEI_Old"] = cei_old
    out["CEI_New"] = cei_new
    out["CEI_Change"] = np.round(cei_new - cei_old, 6)
    out["Band_Old"] = band_label(cei_old)
    out["Band_New"] = band_label(cei_new)
    out["Below_80_Old"] = np.where(low_old, "Yes", "No")
    out["Below_80_New"] = np.where(low_new, "Yes", "No")
    out["Status_Old"] = status(cei_old)
    out["Status_New"] = status(cei_new)
    out["Movement"] = np.select(
        [~low_old & low_new, low_old & ~low_new, low_old & low_new],
        ["Dropped below 80", "Rose to 80+", "Stayed below 80"], "Stayed 80+")
    print(f"{len(out):,} customers, {has.sum():,} with Repair Advice, "
          f"{is_test.sum():,} in the test sample", flush=True)
    return out


# ======================== التنسيق (أبيض وأسود) ========================

FONT = "Arial"
GRID = Side(style="thin", color="BFBFBF")
DARK = Side(style="medium", color="000000")
STYLES = {
    "title": dict(font=Font(name=FONT, size=16, bold=True)),
    "subtitle": dict(font=Font(name=FONT, size=10, italic=True, color="595959")),
    "section": dict(font=Font(name=FONT, size=12, bold=True), border=Border(bottom=DARK)),
    "header": dict(font=Font(name=FONT, size=10, bold=True, color="FFFFFF"),
                   fill=PatternFill("solid", fgColor="000000"),
                   alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)),
    "label": dict(font=Font(name=FONT, size=10), border=Border(bottom=GRID)),
    "bold_label": dict(font=Font(name=FONT, size=10, bold=True), border=Border(top=DARK, bottom=DARK)),
    "text": dict(font=Font(name=FONT, size=10), alignment=Alignment(wrap_text=True, vertical="top")),
    "note": dict(font=Font(name=FONT, size=9, italic=True, color="595959"),
                 alignment=Alignment(wrap_text=True, vertical="top")),
    "shade": dict(font=Font(name=FONT, size=10, bold=True), fill=PatternFill("solid", fgColor="D9D9D9"),
                  border=Border(bottom=GRID)),
}
NUMBER = {"int": "#,##0", "dec": "#,##0.00", "pct": "0.0%", "chg_int": "+#,##0;-#,##0;0",
          "chg_dec": "+#,##0.00;-#,##0.00;0.00", "chg_pct": "+0.0%;-0.0%;0.0%"}


def cell(ws, value, style="label", fmt=None, bold=False):
    c = WriteOnlyCell(ws, value=value)
    for key, val in STYLES[style].items():
        setattr(c, key, val)
    if bold:
        c.font = Font(name=FONT, size=10, bold=True)
        c.border = Border(top=DARK, bottom=DARK)
    if fmt:
        c.number_format = NUMBER.get(fmt, fmt)
        c.alignment = Alignment(horizontal="right")
    return c


def new_sheet(wb, name, widths, title, subtitle):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.append([cell(ws, title, "title")])
    ws.append([cell(ws, subtitle, "subtitle")])
    ws.append([])
    return ws


def header_row(ws, names):
    ws.append([cell(ws, n, "header") for n in names])


# ======================== الشيتات ========================

class Refs:
    """مراجع أعمدة شيت Customer_Data للمعادلات."""

    def __init__(self, columns, n):
        self.n = n
        self.letter = {c: get_column_letter(i) for i, c in enumerate(columns, start=1)}

    def __call__(self, col):
        L = self.letter[col]
        return f"Customer_Data!${L}$2:${L}${self.n + 1}"


def countifs(R, *pairs):
    parts = [f'{R(c)},"{v}"' if isinstance(v, str) and not v.startswith("&") else f"{R(c)},{v[1:]}"
             for c, v in pairs]
    return "COUNTIFS(" + ",".join(parts) + ")"


def write_overview(wb, data, new_weights, weight_source):
    ws = new_sheet(wb, "Overview", [28, 95],
                   "CEI: Old Weights vs New Weights",
                   f"Generated {date.today():%d %B %Y} from {INPUT.name}")
    n, adv, test = len(data), int((data.Has_Repair_Advice == "Yes").sum()), int((data.Sample == "Test").sum())
    rows = [
        ("section", "About this file"),
        ("Input file", INPUT.name),
        ("Weights source", weight_source),
        ("Customers", f"{n:,} ({adv:,} with Repair Advice, {n - adv:,} without)"),
        ("Test sample", f"{test:,} customers the model never saw (random split, seed {SEED}, share {TEST_SHARE:.0%}). "
                        "Official model results are on this sample."),
        ("", ""),
        ("section", "How every number is calculated"),
        ("Achievement", "Customer score ÷ full score of the dimension. Example: Wi-Fi 14 of 20 = 70%."),
        ("New dimension score", "Achievement × new weight. Example: 70% × 39 = 27.30 points."),
        ("CEI", "Sum of the eight dimension scores. Old CEI uses the old weights, New CEI the new weights."),
        ("Below 80", f"A customer is flagged when CEI is below {CEI_THRESHOLD}."),
        ("Status", "Detected = has Repair Advice and CEI < 80.  Hidden = has Repair Advice and CEI ≥ 80.  "
                   "False alarm = no Repair Advice and CEI < 80.  OK = no Repair Advice and CEI ≥ 80."),
        ("Detection rate", "Detected ÷ (Detected + Hidden): share of real problems that CEI exposes."),
        ("False alarm rate", "False alarm ÷ (False alarm + OK): share of healthy customers flagged."),
        ("Precision", "Detected ÷ (Detected + False alarm): of the customers below 80, the share with a real problem."),
        ("", ""),
        ("section", "Sheets"),
        ("Summary", "Headline comparison for all customers and for the test sample."),
        ("Weights", "Old and new weight of each dimension."),
        ("Dimension_Comparison", "Per dimension: weight, average points, and how often it is deducted "
                                 "for customers with and without Repair Advice."),
        ("Band_Distribution", "Customers in each 5-point CEI band, old vs new, with a chart."),
        ("Band_Transition", "Where customers moved: old CEI band (rows) to new CEI band (columns)."),
        ("Status_Transition", "Movement between Detected / Hidden / False alarm / OK, and above/below 80."),
        ("Per_Advice", "For every Repair Advice code: customers, share below 80 and average CEI, old vs new."),
        ("Per_Family", "The same, grouped into five advice families."),
        ("Customer_Data", "One row per customer with old and new scores. Excel table 'CEI_Data', ready for PivotTables."),
        ("", ""),
        ("section", "Building a PivotTable"),
        ("Steps", "Click any cell in Customer_Data → Insert → PivotTable → OK."),
        ("Example 1", "Rows: Band_Old · Columns: Band_New · Values: Count of Serial_Number."),
        ("Example 2", "Rows: Main_Advice · Columns: Status_New · Values: Count of Serial_Number."),
        ("Example 3", "Rows: Advice_Families · Values: Average of CEI_Old and Average of CEI_New · Filter: Sample = Test."),
        ("", ""),
        ("section", "Notes"),
        ("Formulas", "All comparison sheets use COUNTIFS / AVERAGEIFS on Customer_Data, so every figure can be traced. "
                     "Customer_Data itself holds calculated values."),
        ("Advice codes", "A customer can have several codes; Per_Advice counts the customer under each of them, "
                         "so its rows add up to more than the number of customers."),
        ("Weights", "New weights: " + ", ".join(f"{s} {new_weights[c]}" for s, c in zip(SHORT, COLUMNS)) + "."),
    ]
    for key, val in rows:
        if key == "section":
            ws.append([cell(ws, val, "section"), cell(ws, None, "section")])
        elif key == "":
            ws.append([])
        else:
            ws.append([cell(ws, key, "shade"), cell(ws, val, "text")])


def write_summary(wb, R, model_summary):
    ws = new_sheet(wb, "Summary", [44, 16, 16, 14, 62],
                   "Summary: Old Weights vs New Weights",
                   "Every figure is a COUNTIFS / AVERAGEIFS formula on Customer_Data.")
    row = [3]  # عدد الأسطر المكتوبة لحد هلأ (العنوان، الوصف، سطر فاضي)

    def put(cells):
        ws.append(cells)
        row[0] += 1

    def block(title, scope):
        extra = [("Sample", scope)] if scope else []
        sc = f',{R("Sample")},"{scope}"' if scope else ""
        put([cell(ws, title, "section")] + [cell(ws, None, "section")] * 4)
        put([cell(ws, n, "header") for n in ["Measure", "Old weights", "New weights", "Change", "What it means"]])

        def cnt(col, value):
            return "=" + countifs(R, (col, value), *extra)

        def avg(col):
            return f"=AVERAGEIFS({R(col)}{sc})" if scope else f"=AVERAGE({R(col)})"

        metrics = [
            ("customers", "Customers", "=" + countifs(R, ("Sample", scope)) if scope else f"=ROWS({R('CEI_Old')})",
             None, "int", "Customers in this scope."),
            ("avg", "Average CEI", avg("CEI_Old"), avg("CEI_New"), "dec", "Mean CEI of the customers in scope."),
            ("c100", "Customers with CEI = 100", "=" + countifs(R, ("CEI_Old", "&100"), *extra),
             "=" + countifs(R, ("CEI_New", "&100"), *extra), "int", "No deduction in any dimension."),
            ("c90", "Customers with CEI ≥ 90", "=" + countifs(R, ("CEI_Old", ">=90"), *extra),
             "=" + countifs(R, ("CEI_New", ">=90"), *extra), "int", ""),
            ("above", "Customers with CEI ≥ 80", cnt("Below_80_Old", "No"), cnt("Below_80_New", "No"), "int",
             "Not flagged by CEI."),
            ("below", "Customers with CEI < 80", cnt("Below_80_Old", "Yes"), cnt("Below_80_New", "Yes"), "int",
             "Flagged by CEI."),
            ("share", "Share of customers below 80", None, None, "pct", "CEI < 80 ÷ customers."),
            ("det", "Problems detected (advice, CEI < 80)", cnt("Status_Old", "Detected"),
             cnt("Status_New", "Detected"), "int", "Real problems that CEI exposes."),
            ("hid", "Problems hidden (advice, CEI ≥ 80)", cnt("Status_Old", "Hidden"),
             cnt("Status_New", "Hidden"), "int", "Real problems that CEI misses."),
            ("fa", "False alarms (no advice, CEI < 80)", cnt("Status_Old", "False alarm"),
             cnt("Status_New", "False alarm"), "int", "Healthy customers flagged."),
            ("ok", "Healthy and not flagged (OK)", cnt("Status_Old", "OK"), cnt("Status_New", "OK"), "int", ""),
            ("det_r", "Detection rate", None, None, "pct", "Detected ÷ all customers with Repair Advice."),
            ("hid_r", "Hidden rate", None, None, "pct", "Hidden ÷ all customers with Repair Advice."),
            ("fa_r", "False alarm rate", None, None, "pct", "False alarms ÷ all customers without Repair Advice."),
            ("prec", "Precision", None, None, "pct", "Of the customers below 80, the share with Repair Advice."),
        ]
        first = row[0] + 1
        p = {key: first + i for i, (key, *_rest) in enumerate(metrics)}
        ratios = {
            "share": lambda c: f"={c}{p['below']}/{c}{p['customers']}",
            "det_r": lambda c: f"={c}{p['det']}/({c}{p['det']}+{c}{p['hid']})",
            "hid_r": lambda c: f"={c}{p['hid']}/({c}{p['det']}+{c}{p['hid']})",
            "fa_r": lambda c: f"={c}{p['fa']}/({c}{p['fa']}+{c}{p['ok']})",
            "prec": lambda c: f"={c}{p['det']}/({c}{p['det']}+{c}{p['fa']})",
        }
        for key, name, old, new, fmt, note in metrics:
            r = p[key]
            if key in ratios:
                old, new = ratios[key]("B"), ratios[key]("C")
            elif new is None:
                new = f"=B{r}"
            chg = {"int": "chg_int", "dec": "chg_dec", "pct": "chg_pct"}[fmt]
            put([cell(ws, name), cell(ws, old, fmt=fmt), cell(ws, new, fmt=fmt),
                 cell(ws, f"=C{r}-B{r}", fmt=chg), cell(ws, note, "note")])
        put([])

    block("All customers", None)
    block("Test sample (official model results: customers the model never saw)", "Test")

    if model_summary is not None:
        put([cell(ws, "From the model result file (test sample)", "section")] + [cell(ws, None, "section")] * 4)
        put([cell(ws, n, "header") for n in ["Measure", "Old weights", "New weights", "Change", "What it means"]])
        notes = {
            "Ranking_AUC": ("Ranking quality (AUC)", "Chance that a customer with advice gets a lower CEI than one without (1 = perfect)."),
            "Detection_at_Same_False_Alarm_%": ("Detection at the same false-alarm rate",
                                                "Old weights allowed the same false-alarm rate as the new ones: the fair comparison."),
            "Advice_Types_Reaching_80%": ("Advice codes reaching 80% detection",
                                          "Codes with 30+ customers where at least 80% of their customers fall below 80."),
        }
        for measure, (label, note) in notes.items():
            if measure not in model_summary.index:
                continue
            o, f = float(model_summary.loc[measure, "Old"]), float(model_summary.loc[measure, "Final"])
            if measure == "Ranking_AUC":
                fmt, chg = "0.000", "+0.000;-0.000;0.000"
            elif "Types" in measure:
                fmt, chg = "int", "chg_int"
            else:
                fmt, chg, o, f = "pct", "chg_pct", o / 100, f / 100
            put([cell(ws, label), cell(ws, o, fmt=fmt), cell(ws, f, fmt=fmt),
                 cell(ws, f - o, fmt=chg), cell(ws, note, "note")])
        put([cell(ws, f"Source: {RESULT_PATH.name}, sheet Summary (copied values).", "note")])


def write_weights(wb, new_weights):
    ws = new_sheet(wb, "Weights", [22, 14, 14, 12], "Weights", "Points out of 100 for each CEI dimension.")
    header_row(ws, ["Dimension", "Old weight", "New weight", "Change"])
    for i, (s, c) in enumerate(zip(SHORT, COLUMNS)):
        r = 5 + i
        ws.append([cell(ws, s), cell(ws, OLD_WEIGHTS[c], fmt="int"), cell(ws, new_weights[c], fmt="int"),
                   cell(ws, f"=C{r}-B{r}", fmt="chg_int")])
    ws.append([cell(ws, "Total", bold=True), cell(ws, "=SUM(B5:B12)", fmt="int", bold=True),
               cell(ws, "=SUM(C5:C12)", fmt="int", bold=True), cell(ws, "=C13-B13", fmt="chg_int", bold=True)])
    ws.append([])
    ws.append([cell(ws, "Old weights are the full scores of the current CEI. New weights come from the model "
                        "(cei_advice_model.py, sheet Final_Weights).", "note")])


def write_dimensions(wb, R):
    ws = new_sheet(wb, "Dimension_Comparison", [16, 11, 11, 11, 14, 14, 14, 16, 16, 16, 14, 16],
                   "Dimension Comparison",
                   "Why each dimension got its weight: how often it is deducted for customers with and without Repair Advice.")
    header_row(ws, ["Dimension", "Old weight", "New weight", "Change", "Avg points (old)", "Avg points (new)",
                    "Avg achievement", "% deducted: all", "% deducted: with advice", "% deducted: without advice",
                    "Gap (with − without)", "Effect on avg CEI"])
    for i, s in enumerate(SHORT):
        r = 5 + i
        o, nw, old_col = f"Weights!B{5 + i}", f"Weights!C{5 + i}", f"{s}_Old"
        ws.append([
            cell(ws, s),
            cell(ws, f"={o}", fmt="int"), cell(ws, f"={nw}", fmt="int"), cell(ws, f"=C{r}-B{r}", fmt="chg_int"),
            cell(ws, f"=AVERAGE({R(old_col)})", fmt="dec"), cell(ws, f"=AVERAGE({R(s + '_New')})", fmt="dec"),
            cell(ws, f"=E{r}/B{r}", fmt="pct"),
            cell(ws, f'=COUNTIFS({R(old_col)},"<"&B{r})/ROWS({R(old_col)})', fmt="pct"),
            cell(ws, f'=COUNTIFS({R(old_col)},"<"&B{r},{R("Has_Repair_Advice")},"Yes")/COUNTIFS({R("Has_Repair_Advice")},"Yes")', fmt="pct"),
            cell(ws, f'=COUNTIFS({R(old_col)},"<"&B{r},{R("Has_Repair_Advice")},"No")/COUNTIFS({R("Has_Repair_Advice")},"No")', fmt="pct"),
            cell(ws, f"=I{r}-J{r}", fmt="chg_pct"),
            cell(ws, f"=F{r}-E{r}", fmt="chg_dec"),
        ])
    ws.append([cell(ws, "Total", bold=True), cell(ws, "=SUM(B5:B12)", fmt="int", bold=True),
               cell(ws, "=SUM(C5:C12)", fmt="int", bold=True), cell(ws, "=SUM(D5:D12)", fmt="chg_int", bold=True),
               cell(ws, "=SUM(E5:E12)", fmt="dec", bold=True), cell(ws, "=SUM(F5:F12)", fmt="dec", bold=True)]
              + [cell(ws, None, bold=True)] * 5 + [cell(ws, "=SUM(L5:L12)", fmt="chg_dec", bold=True)])
    ws.append([])
    for note in [
        "Avg points: average score of the dimension across all customers. Their total is the average CEI.",
        "% deducted: share of customers who lost at least one point in the dimension.",
        "Gap: a large positive gap means deductions in this dimension point to real problems, so the model raised its weight. "
        "A small or negative gap (Service, Stability) means deductions there do not separate problem customers, so the weight went to the minimum.",
        "Effect on avg CEI: how much the dimension moved the average CEI (new points minus old points).",
    ]:
        ws.append([cell(ws, note, "note")])
    ws.freeze_panes = "B5"


def write_bands(wb, R, bands):
    ws = new_sheet(wb, "Band_Distribution", [14, 13, 13, 12, 13, 13, 13, 13, 11, 11],
                   "CEI Band Distribution", "Number of customers in each 5-point CEI band.")
    header_row(ws, ["CEI band", "All: old", "All: new", "Change", "With advice: old", "With advice: new",
                    "Without advice: old", "Without advice: new", "% of all: old", "% of all: new"])
    first = 5
    for i, b in enumerate(bands):
        r = first + i
        ws.append([
            cell(ws, b),
            cell(ws, "=" + countifs(R, ("Band_Old", b)), fmt="int"),
            cell(ws, "=" + countifs(R, ("Band_New", b)), fmt="int"),
            cell(ws, f"=C{r}-B{r}", fmt="chg_int"),
            cell(ws, "=" + countifs(R, ("Band_Old", b), ("Has_Repair_Advice", "Yes")), fmt="int"),
            cell(ws, "=" + countifs(R, ("Band_New", b), ("Has_Repair_Advice", "Yes")), fmt="int"),
            cell(ws, "=" + countifs(R, ("Band_Old", b), ("Has_Repair_Advice", "No")), fmt="int"),
            cell(ws, "=" + countifs(R, ("Band_New", b), ("Has_Repair_Advice", "No")), fmt="int"),
            cell(ws, f"=B{r}/B${first + len(bands)}", fmt="pct"),
            cell(ws, f"=C{r}/C${first + len(bands)}", fmt="pct"),
        ])
    last = first + len(bands) - 1
    ws.append([cell(ws, "Total", bold=True)] +
              [cell(ws, f"=SUM({L}{first}:{L}{last})", fmt="chg_int" if L == "D" else ("pct" if L in "IJ" else "int"),
                    bold=True) for L in "BCDEFGHIJ"])
    ws.append([])
    ws.append([cell(ws, "Bands include the lower edge and exclude the upper one (80-85 means 80 ≤ CEI < 85); "
                        "CEI = 100 is in 95-100.", "note")])

    chart = BarChart()
    chart.type = "col"
    chart.title = "Customers per CEI band: old vs new weights"
    chart.y_axis.title = "Customers"
    chart.x_axis.title = "CEI band"
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=4, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    for series, colour in zip(chart.series, ["A6A6A6", "000000"]):
        series.graphicalProperties.solidFill = colour
        series.graphicalProperties.line.solidFill = colour
    chart.y_axis.majorGridlines = None
    chart.y_axis.delete = False
    chart.x_axis.delete = False
    chart.height, chart.width = 9, 24
    ws.add_chart(chart, f"A{last + 6}")


def write_transition(wb, R):
    edges = [(0, 60, "Below 60"), (60, 70, "60-70"), (70, 80, "70-80"), (80, 90, "80-90"), (90, 100.001, "90-100")]
    ws = new_sheet(wb, "Band_Transition", [26] + [13] * 6,
                   "Band Transition", "Rows: old CEI band. Columns: new CEI band. Cells: number of customers.")

    def matrix(title, extra):
        ws.append([cell(ws, title, "section")] + [cell(ws, None, "section")] * 6)
        header_row(ws, ["Old band ↓ / New band →"] + [e[2] for e in edges] + ["Total"])
        top = pos[0] + 2
        for i, (lo, hi, name) in enumerate(edges):
            r = top + i
            row = [cell(ws, name, "shade")]
            for lo2, hi2, _ in edges:
                f = countifs(R, ("CEI_Old", f"&\">=\"&{lo}"), ("CEI_Old", f"&\"<\"&{hi}"),
                             ("CEI_New", f"&\">=\"&{lo2}"), ("CEI_New", f"&\"<\"&{hi2}"), *extra)
                row.append(cell(ws, "=" + f, fmt="int"))
            row.append(cell(ws, f"=SUM(B{r}:F{r})", fmt="int", bold=True))
            ws.append(row)
        last = top + len(edges) - 1
        ws.append([cell(ws, "Total", bold=True)] +
                  [cell(ws, f"=SUM({L}{top}:{L}{last})", fmt="int", bold=True) for L in "BCDEFG"])
        ws.append([])
        pos[0] += len(edges) + 4

    pos = [3]
    matrix("All customers", [])
    matrix("Customers with Repair Advice", [("Has_Repair_Advice", "Yes")])
    matrix("Customers without Repair Advice", [("Has_Repair_Advice", "No")])
    ws.append([cell(ws, "Cells below the diagonal: CEI fell to a lower band. Above the diagonal: CEI rose.", "note")])


def write_status(wb, R):
    states = ["Detected", "Hidden", "False alarm", "OK"]
    moves = ["Dropped below 80", "Rose to 80+", "Stayed below 80", "Stayed 80+"]
    ws = new_sheet(wb, "Status_Transition", [30] + [14] * 5,
                   "Status Transition", "How each customer's status changed from the old to the new weights.")
    pos = [3]

    def matrix(title, extra):
        ws.append([cell(ws, title, "section")] + [cell(ws, None, "section")] * 5)
        header_row(ws, ["Old status ↓ / New status →"] + states + ["Total"])
        top = pos[0] + 2
        for i, s in enumerate(states):
            r = top + i
            ws.append([cell(ws, s, "shade")] +
                      [cell(ws, "=" + countifs(R, ("Status_Old", s), ("Status_New", t), *extra), fmt="int")
                       for t in states] + [cell(ws, f"=SUM(B{r}:E{r})", fmt="int", bold=True)])
        last = top + len(states) - 1
        ws.append([cell(ws, "Total", bold=True)] +
                  [cell(ws, f"=SUM({L}{top}:{L}{last})", fmt="int", bold=True) for L in "BCDEF"])
        ws.append([])
        pos[0] += len(states) + 4

    matrix("All customers", [])
    matrix("Test sample", [("Sample", "Test")])

    ws.append([cell(ws, "Movement across the 80 line", "section")] + [cell(ws, None, "section")] * 5)
    header_row(ws, ["Movement", "All customers", "With advice", "Without advice", "Share of all", ""])
    top = pos[0] + 2
    for i, m in enumerate(moves):
        r = top + i
        ws.append([cell(ws, m, "shade"),
                   cell(ws, "=" + countifs(R, ("Movement", m)), fmt="int"),
                   cell(ws, "=" + countifs(R, ("Movement", m), ("Has_Repair_Advice", "Yes")), fmt="int"),
                   cell(ws, "=" + countifs(R, ("Movement", m), ("Has_Repair_Advice", "No")), fmt="int"),
                   cell(ws, f"=B{r}/SUM(B${top}:B${top + 3})", fmt="pct")])
    ws.append([cell(ws, "Total", bold=True)] +
              [cell(ws, f"=SUM({L}{top}:{L}{top + 3})", fmt="int", bold=True) for L in "BCD"] +
              [cell(ws, f"=SUM(E{top}:E{top + 3})", fmt="pct", bold=True)])
    ws.append([])
    ws.append([cell(ws, "Hidden = has Repair Advice but CEI ≥ 80. False alarm = no Repair Advice but CEI < 80.", "note")])


def write_groups(wb, R, name, title, items, column, key_fn, label_header, extra_col=None):
    widths = [52, 18, 12, 12, 13, 13, 12, 12, 12, 12, 12, 13, 13]
    ws = new_sheet(wb, name, widths, title,
                   "Share below 80 = customers with CEI < 80 ÷ customers in the group. "
                   "Test columns use only the test sample (official results).")
    header_row(ws, [label_header, "Family" if extra_col else "", "Customers", "Test customers",
                    "Below 80: old", "Below 80: new", "Change", "Test: old", "Test: new",
                    "Hidden: old", "Hidden: new", "Avg CEI: old", "Avg CEI: new"])
    for i, item in enumerate(items):
        r = 5 + i
        crit = f'"*{key_fn(item)}*"'
        base = f"{R(column)},{crit}"
        test = f'{base},{R("Sample")},"Test"'
        ws.append([
            cell(ws, item),
            cell(ws, extra_col(item) if extra_col else None),
            cell(ws, f"=COUNTIFS({base})", fmt="int"),
            cell(ws, f"=COUNTIFS({test})", fmt="int"),
            cell(ws, f'=COUNTIFS({base},{R("Below_80_Old")},"Yes")/C{r}', fmt="pct"),
            cell(ws, f'=COUNTIFS({base},{R("Below_80_New")},"Yes")/C{r}', fmt="pct"),
            cell(ws, f"=F{r}-E{r}", fmt="chg_pct"),
            cell(ws, f'=IFERROR(COUNTIFS({test},{R("Below_80_Old")},"Yes")/D{r},"")', fmt="pct"),
            cell(ws, f'=IFERROR(COUNTIFS({test},{R("Below_80_New")},"Yes")/D{r},"")', fmt="pct"),
            cell(ws, f'=COUNTIFS({base},{R("Below_80_Old")},"No")', fmt="int"),
            cell(ws, f'=COUNTIFS({base},{R("Below_80_New")},"No")', fmt="int"),
            cell(ws, f"=AVERAGEIFS({R('CEI_Old')},{base})", fmt="dec"),
            cell(ws, f"=AVERAGEIFS({R('CEI_New')},{base})", fmt="dec"),
        ])
    ws.append([])
    ws.append([cell(ws, "A customer with several advice codes is counted under each code, so the rows overlap. "
                        "Below 80 = the share of the group that CEI flags (detection).", "note")])
    ws.freeze_panes = "B5"


def write_customer_data(wb, data):
    print("Writing Customer_Data...", flush=True)
    ws = wb.create_sheet("Customer_Data")
    widths = {"Serial_Number": 20, "Advice_Codes": 48, "Main_Advice": 40, "Advice_Families": 26}
    for i, c in enumerate(data.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 13)
    ws.freeze_panes = "B2"
    header_row(ws, list(data.columns))
    decimal_cols = {i for i, c in enumerate(data.columns)
                    if c.endswith("_New") and c not in ("Band_New", "Below_80_New", "Status_New")
                    or c in ("CEI_New", "CEI_Change")}
    fmt_cells = {}
    for row in data.itertuples(index=False):
        values = list(row)
        for i in decimal_cols:
            c = WriteOnlyCell(ws, value=float(values[i]))
            c.number_format = "0.00"
            values[i] = c
        ws.append(values)
    ref = f"A1:{get_column_letter(len(data.columns))}{len(data) + 1}"
    table = Table(displayName="CEI_Data", ref=ref)
    table.tableColumns = [TableColumn(id=i, name=c) for i, c in enumerate(data.columns, start=1)]
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)


def main():
    new_weights, weight_source = load_weights()
    data = build_data(new_weights)
    model_summary = load_model_summary()
    R = Refs(list(data.columns), len(data))

    bands = sorted(set(data["Band_Old"]) | set(data["Band_New"]))
    codes = (data.loc[data.Advice_Count > 0, "Advice_Codes"].str.split(", ").explode().value_counts())

    wb = Workbook(write_only=True)
    write_overview(wb, data, new_weights, weight_source)
    write_summary(wb, R, model_summary)
    write_weights(wb, new_weights)
    write_dimensions(wb, R)
    write_bands(wb, R, bands)
    write_transition(wb, R)
    write_status(wb, R)
    write_groups(wb, R, "Per_Advice", "Per Repair Advice code", list(codes.index), "Advice_Codes",
                 lambda c: c, "Repair Advice code", extra_col=advice_family)
    families = [f for f in FAMILIES if (data.Advice_Families.str.contains(f, regex=False)).any()]
    write_groups(wb, R, "Per_Family", "Per advice family", families, "Advice_Families",
                 lambda f: f, "Advice family")
    write_customer_data(wb, data)
    wb.calculation.fullCalcOnLoad = True  # Excel بيحسب المعادلات أول ما يفتح الملف
    print("Saving...", flush=True)
    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:  # python cei_apply_weights.py <merged.csv> [result.xlsx] [output.xlsx]
        FILE_PATH = sys.argv[1]
        INPUT = Path(FILE_PATH)
        RESULT_PATH = Path(sys.argv[2]) if len(sys.argv) > 2 else INPUT.with_name(INPUT.stem + "_advice_model_result.xlsx")
        OUTPUT_PATH = Path(sys.argv[3]) if len(sys.argv) > 3 else INPUT.with_name(INPUT.stem + "_new_weights_applied.xlsx")
    main()
